from datetime import date
from django.db.models import Sum,Max,Q
from django.utils import json, timezone
from datetime import timedelta
from django.shortcuts import render, redirect,get_object_or_404
from django.contrib import messages
from Account.models import ContactInfo, Nurse,Staff,Student,StudentVisit, MedicineInventory
from django.http import HttpResponse
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas
from datetime import datetime
import django.utils.json as json
from django.db import transaction
from django.http import HttpResponse
import json
def nurse_dashboard(request):

    # Check Login
    if 'nurse_id' not in request.session:
        return redirect('nurse_login')

    # Logged-in Nurse
    nurse = Nurse.objects.get(
        nurse_id=request.session['nurse_id']
    )

    today = date.today()

    # Dashboard Statistics
    todays_visits_count = StudentVisit.objects.filter(
        date=today
    ).count()

    total_patients = StudentVisit.objects.values(
        'reg_no'
    ).distinct().count()

    chronic_patients_count = StudentVisit.objects.filter(
        chronic_illness="Yes"
    ).values(
        'reg_no'
    ).distinct().count()

    low_stock_count = MedicineInventory.objects.filter(
        quantity__lte=10
    ).count()

    # Recent Visits
    recent_visits = StudentVisit.objects.order_by(
        '-date',
        '-time'
    )[:10]

    context = {

        "page": "dashboard",

        "nurse": nurse,

        "todays_visits_count": todays_visits_count,

        "total_patients": total_patients,

        "chronic_patients_count": chronic_patients_count,

        "low_stock_count": low_stock_count,

        "recent_visits": recent_visits,

    }

    return render(
        request,
        "Nurse/nurse_dashboard.html",
        context
    )
def student_visit(request, reg_no):
    student = get_object_or_404(Student, reg_no=reg_no)

    if request.method == "POST":
        medicines = request.POST.getlist("medicine[]")
        dosages = request.POST.getlist("dosage[]")
        card_finished = request.POST.getlist("card_finished[]")

        medicine_details = [
            f"{medicine} - {dosage}"
            for medicine, dosage in zip(medicines, dosages)
            if medicine and dosage
        ]

        medicine_text = ", ".join(medicine_details)

        with transaction.atomic():
            StudentVisit.objects.create(
                student=student,
                reg_no=student.reg_no,
                student_name=student.name,
                department=student.department,
                blood_group=student.blood_group,
                problem=request.POST.get("problem"),
                medicine=medicine_text,
                staff_name=request.POST.get("staff_name") or None,
                chronic_illness=request.POST.get("chronic_illness", "No"),
                remarks=request.POST.get("remarks") or None,
                follow_up_days=request.POST.get("follow_up_days") or None,
            )

            for medicine_name in card_finished:
                if medicine_name:
                    inventory = MedicineInventory.objects.select_for_update().filter(
                        medicine_name=medicine_name
                    ).first()

                    if inventory and inventory.quantity > 0:
                        inventory.quantity -= 1
                        inventory.save(update_fields=["quantity"])

        messages.success(request, "Student visit saved successfully.")
        return redirect("student_search")

    medicines = MedicineInventory.objects.all().order_by("medicine_name")
    departments = Staff.objects.values("department").distinct().order_by("department")
    staffs = Staff.objects.all().order_by("department", "name")

    context = {
        "student": student,
        "medicines": medicines,
        "departments": departments,
        "staffs": staffs,
        "problems": StudentVisit.PROBLEM_CHOICES,
        "chronic_choices": StudentVisit.CHRONIC_CHOICES,
        "page": "student_visit",
    }

    return render(request, "Nurse/student_visit.html", context)
def student_search(request):

    reg_no = request.GET.get("reg_no")

    if reg_no:
        return redirect("student_visit", reg_no=reg_no)

    return render(request, "Nurse/student_search.html")

from django.db.models import Q
def visit_history(request):
    visits = StudentVisit.objects.select_related("student").all().order_by("-date", "-time")

    search_date = request.GET.get("date", "").strip()
    problem = request.GET.get("problem", "").strip()
    reg_no = request.GET.get("reg_no", "").strip()

    if search_date:
        visits = visits.filter(date=search_date)
    if problem:
        visits = visits.filter(problem=problem)
    if reg_no:
        visits = visits.filter(student__reg_no__icontains=reg_no)

    context = {
        "visits": visits,
        "problems": StudentVisit.PROBLEM_CHOICES,
        "selected_problem": problem,
        "selected_date": search_date,
        "reg_no": reg_no,
        "page": "visit_history",
    }

    return render(request, "Nurse/visit_history.html", context)
def inventory(request):

    if request.method == "POST":

        action = request.POST.get("action")

        if action == "add_medicine":

            medicine_name = request.POST.get("medicine_name", "").strip()
            quantity = request.POST.get("quantity", 0)
            expiry_date = request.POST.get("expiry_date")
            description = request.POST.get("description", "").strip()

            if MedicineInventory.objects.filter(medicine_name__iexact=medicine_name).exists():

                messages.error(request, "Medicine already exists.")

            else:

                MedicineInventory.objects.create(
                    medicine_name=medicine_name,
                    quantity=int(quantity),
                    expiry_date=expiry_date,
                    description=description
                )

                messages.success(request, "Medicine added successfully.")


        elif action == "update_stock":

            medicine_id = request.POST.get("medicine_id")
            add_quantity = request.POST.get("add_quantity")
            expiry_date = request.POST.get("expiry_date")

            medicine = MedicineInventory.objects.filter(medicine_id=medicine_id).first()

            if medicine:

                medicine.quantity += int(add_quantity)
                medicine.expiry_date = expiry_date
                medicine.save()

                messages.success(request, "Stock updated successfully.")

            else:

                messages.error(request, "Medicine not found.")


        return redirect("inventory")


    search = request.GET.get("search", "").strip()

    medicines = MedicineInventory.objects.all().order_by("medicine_name")

    if search:

        medicines = medicines.filter(
            medicine_name__icontains=search
        )


    total_medicines = medicines.count()

    total_stock = sum(
        medicine.quantity
        for medicine in medicines
    )

    low_stock = medicines.filter(
        quantity__gt=0,
        quantity__lte=3
    ).count()


    context = {
        "medicines": medicines,
        "total_medicines": total_medicines,
        "total_stock": total_stock,
        "low_stock": low_stock,
        "search": search
    }

    return render(
        request,
        "Nurse/inventory.html",
        context
    )
def chronic_illness(request):

    # -----------------------------
    # Update Follow-up & Remarks
    # -----------------------------
    if request.method == "POST":

        visit_id = request.POST.get("visit_id")

        try:
            visit = StudentVisit.objects.get(visit_id=visit_id)

            follow_up = request.POST.get("follow_up_days")

            if follow_up:
                visit.follow_up_days = int(follow_up)
            else:
                visit.follow_up_days = None

            visit.remarks = request.POST.get("remarks", "")

            visit.save()

            messages.success(
                request,
                "Follow-up details updated successfully."
            )

        except StudentVisit.DoesNotExist:

            messages.error(
                request,
                "Patient record not found."
            )

        return redirect("chronic_illness")

    # -----------------------------
    # Search
    # -----------------------------
    search = request.GET.get("search", "")

    latest_visits = (
        StudentVisit.objects
        .filter(chronic_illness="Yes")
        .values("reg_no")
        .annotate(latest_visit=Max("visit_id"))
    )

    latest_ids = [
        item["latest_visit"]
        for item in latest_visits
    ]

    visits = StudentVisit.objects.filter(
        visit_id__in=latest_ids
    ).order_by("-date", "-time")

    if search:

        visits = visits.filter(

            Q(reg_no__icontains=search) |
            Q(student_name__icontains=search)

        )

    today = timezone.now().date()

    total_patients = visits.count()

    upcoming = 0
    due_today = 0
    overdue = 0

    # -----------------------------
    # Calculate Follow-up Status
    # -----------------------------
    for visit in visits:

        visit.next_checkup = None
        visit.status = "No Follow-up"
        visit.message = "No Follow-up Required"

        if visit.follow_up_days:

            visit.next_checkup = (
                visit.date +
                timedelta(days=visit.follow_up_days)
            )

            remaining = (
                visit.next_checkup -
                today
            ).days

            if remaining > 0:

                visit.status = "Upcoming"

                visit.message = (
                    f"{remaining} day(s) remaining"
                )

                upcoming += 1

            elif remaining == 0:

                visit.status = "Today"

                visit.message = (
                    "Check-up Due Today"
                )

                due_today += 1

            else:

                visit.status = "Overdue"

                visit.message = (
                    f"{abs(remaining)} day(s) overdue"
                )

                overdue += 1

    context = {

        "visits": visits,

        "search": search,

        "total_patients": total_patients,

        "upcoming": upcoming,

        "due_today": due_today,

        "overdue": overdue,

        "page": "chronic_illness",

    }

    return render(
        request,
        "Nurse/chronic_illness.html",
        context
    )
def inventory_request(request):

    inventory = MedicineInventory.objects.all().order_by("medicine_name")

    context = {
        "inventory": inventory
    }

    return render(
        request,
        "Nurse/inventory_request.html",
        context
    )
def download_inventory_pdf(request):

    if request.method != "POST":
        return HttpResponse("Invalid Request")

    items_json = request.POST.get("items", "")
    remarks = request.POST.get("remarks", "").strip()

    if not items_json:
        return HttpResponse("No items selected.")

    try:
        items = json.loads(items_json)
    except json.JSONDecodeError:
        return HttpResponse("Invalid item data.")

    if not items:
        return HttpResponse("No items selected.")

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="Inventory_Purchase_Request.pdf"'

    pdf = canvas.Canvas(response)

    width, height = pdf._pagesize
    y = height - 50

    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawCentredString(width / 2, y, "SHC STUDENT HEALTH CARE SYSTEM")

    y -= 20
    pdf.setFont("Helvetica", 12)
    pdf.drawCentredString(width / 2, y, "Sacred Heart College (Autonomous)")

    y -= 18
    pdf.drawCentredString(width / 2, y, "Tirupattur - 635601")

    y -= 30
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawCentredString(width / 2, y, "INVENTORY PURCHASE REQUEST")

    y -= 20
    pdf.line(40, y, 560, y)

    y -= 25
    pdf.setFont("Helvetica", 11)
    pdf.drawString(50, y, f"Date : {datetime.now().strftime('%d-%m-%Y %I:%M %p')}")

    y -= 30

    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(50, y, "No")
    pdf.drawString(100, y, "Item")
    pdf.drawString(430, y, "Quantity")

    y -= 10
    pdf.line(40, y, 560, y)

    y -= 20
    pdf.setFont("Helvetica", 11)

    for index, item in enumerate(items, start=1):

        item_name = str(item.get("item", ""))
        quantity = str(item.get("qty", ""))

        pdf.drawString(50, y, str(index))
        pdf.drawString(100, y, item_name[:50])
        pdf.drawString(450, y, quantity)

        y -= 20

        if y < 100:

            pdf.showPage()

            width, height = pdf._pagesize
            y = height - 50

            pdf.setFont("Helvetica-Bold", 11)
            pdf.drawString(50, y, "No")
            pdf.drawString(100, y, "Item")
            pdf.drawString(430, y, "Quantity")

            y -= 10
            pdf.line(40, y, 560, y)

            y -= 20
            pdf.setFont("Helvetica", 11)

    y -= 20

    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(50, y, "Remarks")

    y -= 18

    pdf.setFont("Helvetica", 11)
    pdf.drawString(50, y, remarks if remarks else "Stock refill for Student Health Care Centre.")

    y -= 70

    pdf.line(60, y, 220, y)
    pdf.line(350, y, 510, y)

    pdf.drawString(95, y - 15, "Requested By")
    pdf.drawString(390, y - 15, "Approved By")

    pdf.setFont("Helvetica-Oblique", 9)
    pdf.drawCentredString(width / 2, 30, "Generated by SHC Student Health Care System")

    pdf.save()

    return response
def contact_enquiries(request):

    if request.method == "POST":

        enquiry_id = request.POST.get("enquiry_id")

        enquiry = ContactInfo.objects.filter(id=enquiry_id).first()

        if enquiry:
            enquiry.delete()

        return redirect("contact_enquiries")

    enquiries = ContactInfo.objects.all().order_by("-id")

    return render(request, "Nurse/contact_enquiries.html", {"enquiries": enquiries})
def nurse_logout(request):
    request.session.flush()
    storage = messages.get_messages(request)

    for _ in storage:
        pass

    return redirect("home")
from openpyxl import Workbook

def download_inventory_excel(request):
    search = request.GET.get("search", "").strip()

    medicines = MedicineInventory.objects.all().order_by("medicine_name")
    if search:
        medicines = medicines.filter(medicine_name__icontains=search)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Medicine Inventory"

    sheet.append([
        "S.No",
        "Medicine Name",
        "Available Stock",
        "Expiry Date",
        "Status",
        "Description",
    ])

    for index, medicine in enumerate(medicines, start=1):
        if medicine.quantity == 0:
            status = "Out of Stock"
        elif medicine.quantity <= 3:
            status = "Low Stock"
        else:
            status = "In Stock"

        sheet.append([
            index,
            medicine.medicine_name,
            medicine.quantity,
            medicine.expiry_date.strftime("%d-%m-%Y"),
            status,
            medicine.description or "",
        ])

    for column in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max_length + 3, 40)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="SHC_Medicine_Inventory.xlsx"'

    workbook.save(response)
    return response

from openpyxl import Workbook
from django.http import HttpResponse
from Account.models import StudentVisit


def download_visit_history_excel(request):
    visits = StudentVisit.objects.using("default").all().order_by("-date", "-time")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Student Visit History"

    sheet.append([
        "Visit ID", "Register Number", "Student Name",
        "Department", "Blood Group", "Problem", "Medicine",
        "Staff Name", "Remarks", "Chronic Illness",
        "Follow-up Days", "Date", "Time"
    ])

    for visit in visits.iterator():
        sheet.append([
            visit.visit_id,
            visit.reg_no,
            visit.student_name,
            visit.department,
            visit.blood_group or "",
            visit.problem,
            visit.medicine,
            visit.staff_name or "",
            visit.remarks or "",
            visit.chronic_illness,
            visit.follow_up_days if visit.follow_up_days is not None else "",
            visit.date.strftime("%d-%m-%Y") if visit.date else "",
            visit.time.strftime("%I:%M:%S %p") if visit.time else "",
        ])

    for column in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max_length + 3, 40)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="SHC_Student_Visit_History.xlsx"'

    workbook.save(response)
    return response

def download_chronic_illness_excel(request):
    search = request.GET.get("search", "").strip()

    # Export only chronic illness records
    visits = StudentVisit.objects.filter(
        chronic_illness="Yes"
    ).order_by("-date", "-time")

    # Apply the same search filter as the page
    if search:
        from django.db.models import Q

        visits = visits.filter(
            Q(reg_no__icontains=search) |
            Q(student_name__icontains=search)
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Chronic Illness"

    # Excel column headings
    sheet.append([
        "Visit ID",
        "Register Number",
        "Student Name",
        "Department",
        "Blood Group",
        "Problem",
        "Medicine",
        "Last Visit",
        "Follow-up Days",
        "Next Check-up",
        "Remarks",
        "Staff Name",
        "Chronic Illness",
    ])

    for visit in visits.iterator():
        next_checkup = ""

        if visit.follow_up_days is not None and visit.date:
            from datetime import timedelta
            next_checkup = visit.date + timedelta(
                days=visit.follow_up_days
            )

        sheet.append([
            visit.visit_id,
            visit.reg_no,
            visit.student_name,
            visit.department,
            visit.blood_group or "",
            visit.problem,
            visit.medicine,
            visit.date.strftime("%d-%m-%Y") if visit.date else "",
            visit.follow_up_days if visit.follow_up_days is not None else "",
            next_checkup.strftime("%d-%m-%Y") if next_checkup else "",
            visit.remarks or "",
            visit.staff_name or "",
            visit.chronic_illness,
        ])

    # Style the header
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="DC3545"
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    # Automatically size columns
    for column in sheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column
        )
        sheet.column_dimensions[
            get_column_letter(column[0].column)
        ].width = min(max_length + 3, 40)

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )
    response["Content-Disposition"] = (
        'attachment; filename="SHC_Chronic_Illness_Followup.xlsx"'
    )

    workbook.save(response)
    return response
